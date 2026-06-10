from __future__ import annotations

import argparse
import json
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter


DEFAULT_HEADERS = ["分组", "标题", "等级", "前置条件", "步骤", "预期结果", "类型", "备注"]
DEFAULT_OUTPUT = "output/generated_cases.xlsx"


def load_cases(path: Path) -> list[dict[str, Any]]:
    payload = json.loads(path.read_text(encoding="utf-8"))
    if isinstance(payload, list):
        return [_normalize_case(item) for item in payload]
    if isinstance(payload, dict) and isinstance(payload.get("cases"), list):
        return [_normalize_case(item) for item in payload["cases"]]
    raise ValueError("Input JSON must be a list of cases or an object with a cases list")


def _normalize_case(item: Any) -> dict[str, Any]:
    if not isinstance(item, dict):
        raise ValueError("Each case must be an object")
    normalized = {header: _stringify_cell(item.get(header, "")) for header in DEFAULT_HEADERS}
    for key, value in item.items():
        if key not in normalized:
            normalized[key] = value
    return normalized


def _stringify_cell(value: Any) -> str:
    if value is None:
        return ""
    if isinstance(value, list):
        return "\n".join(str(part) for part in value)
    return str(value)


def export_cases(cases: list[dict[str, Any]], output_path: Path) -> None:
    output_path.parent.mkdir(parents=True, exist_ok=True)

    workbook = Workbook()
    sheet = workbook.active
    sheet.title = "测试用例"

    sheet.append(DEFAULT_HEADERS)
    for case in cases:
        sheet.append([case.get(header, "") for header in DEFAULT_HEADERS])

    header_fill = PatternFill("solid", fgColor="D9EAF7")
    header_font = Font(bold=True)
    wrap_top = Alignment(wrap_text=True, vertical="top")

    for cell in sheet[1]:
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = Alignment(horizontal="center", vertical="center")

    for row in sheet.iter_rows(min_row=2):
        for cell in row:
            cell.alignment = wrap_top

    widths = {
        "分组": 26,
        "标题": 34,
        "等级": 10,
        "前置条件": 42,
        "步骤": 48,
        "预期结果": 52,
        "类型": 14,
        "备注": 28,
    }
    for index, header in enumerate(DEFAULT_HEADERS, start=1):
        sheet.column_dimensions[get_column_letter(index)].width = widths[header]

    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    workbook.save(output_path)


def build_parser() -> argparse.ArgumentParser:
    parser = argparse.ArgumentParser(description="Export generated test cases to a local Excel file only.")
    parser.add_argument("--input", required=True, type=Path, help="Generated cases JSON file. Excel is not accepted as input.")
    parser.add_argument("--output", default=DEFAULT_OUTPUT, type=Path, help=f"Output .xlsx path, default: {DEFAULT_OUTPUT}")
    return parser


def main() -> None:
    args = build_parser().parse_args()
    if args.input.suffix.lower() in {".xlsx", ".xls"}:
        raise ValueError("Excel is output-only in this workflow; provide generated cases as JSON input")
    if args.output.suffix.lower() != ".xlsx":
        raise ValueError("Output path must end with .xlsx")

    cases = load_cases(args.input)
    export_cases(cases, args.output)
    print(f"Exported {len(cases)} cases to {args.output}")


if __name__ == "__main__":
    main()
